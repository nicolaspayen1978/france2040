#!/usr/bin/env python3
"""Phase 2.1 — use-of-funds incentives only.

€700bn origination fixed (same as Central). Household cash and outstanding unchanged.
Does not wire imports, labour, inflation, or bank losses.
Not a ranking. Not a forecast. Does not close any EC.
"""

from __future__ import annotations

from copy import deepcopy
from pathlib import Path
import sys

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(Path(__file__).resolve().parent))
from modelPhase2 import ALLOC, run, row2040, sum_field

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "Docs" / "21_Model_Phase2.1.xlsx"

BASE = ALLOC["Diversifié-1"]
USES = ["Trav", "Conso", "Inv", "Exist", "Liq", "Sub", "Trf"]


def assert_unit(alloc: dict, name: str) -> dict:
    s = sum(alloc[k] for k in USES)
    if abs(s - 1.0) > 1e-9:
        raise SystemExit(f"{name} shares sum to {s}, not 1")
    return alloc


def shift_from(base: dict, dest: str, sources: list[str], pp: float) -> dict:
    """Move `pp` percentage points equally from sources into dest (Σ)."""
    p = pp / 100.0
    out = {k: float(base[k]) for k in USES}
    share = p / len(sources)
    taken = 0.0
    for src in sources:
        take = min(out[src], share)
        out[src] -= take
        taken += take
    leftover = p - taken
    for src in sources:
        if leftover <= 1e-15:
            break
        take = min(out[src], leftover)
        out[src] -= take
        leftover -= take
        taken += take
    out[dest] += taken
    return assert_unit(out, f"{dest}+{pp}")


def reduce_exist_into_trav_inv(base: dict, pp: float) -> dict:
    """Anti-asset-rotation: Exist down, Trav/Inv up in Diversifié-1 Trav:Inv ratio (Σ)."""
    p = pp / 100.0
    out = {k: float(base[k]) for k in USES}
    take = min(out["Exist"], p)
    out["Exist"] -= take
    t, i = BASE["Trav"], BASE["Inv"]
    out["Trav"] += take * t / (t + i)
    out["Inv"] += take * i / (t + i)
    return assert_unit(out, f"Exist-{pp}")


def additional_spend(alloc: dict) -> float:
    return alloc["Trav"] + alloc["Conso"] + alloc["Inv"]


def metrics(x: dict, r: dict) -> dict:
    x40 = row2040(x)
    r40 = row2040(r)
    gross = x40["cumul"]
    vol = sum_field(x, "vol") - sum_field(r, "vol")
    rec = sum_field(x, "pact_rec") - sum_field(r, "pact_rec")
    return dict(
        name=x["traj"],
        vol=vol,
        rec=rec,
        d_ratio_pp=100 * (x40["ratio"] - r40["ratio"]),
        gross=gross,
        outstanding=x40["outstanding"],
        released=sum_field(x, "released"),
        vol_per_euro=vol / gross if gross else None,
        rec_per_euro=rec / gross if gross else None,
        spend=gross * additional_spend(x["alloc"]) if isinstance(x["alloc"], dict) else None,
    )


def write_sheet(wb, title: str, headers: list[str], rows: list[list]):
    ws = wb.create_sheet(title)
    for i, h in enumerate(headers, 1):
        c = ws.cell(1, i, h)
        c.font = Font(bold=True)
    for r, row in enumerate(rows, 2):
        for i, v in enumerate(row, 1):
            cell = ws.cell(r, i, v)
            if isinstance(v, float):
                cell.number_format = "0.00"
    for i in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(i)].width = 18
    return ws


def main() -> None:
    r = run("R")
    c = run("C", label="C")
    w = run("W", label="W")

    cases = {
        "C": run("C", alloc_override=deepcopy(BASE), label="C"),
        "reno_5": run("C", alloc_override=shift_from(BASE, "Trav", ["Liq", "Exist"], 5), label="reno_5"),
        "reno_10": run("C", alloc_override=shift_from(BASE, "Trav", ["Liq", "Exist"], 10), label="reno_10"),
        "reno_15": run("C", alloc_override=shift_from(BASE, "Trav", ["Liq", "Exist"], 15), label="reno_15"),
        "inv_5": run("C", alloc_override=shift_from(BASE, "Inv", ["Liq", "Exist"], 5), label="inv_5"),
        "inv_10": run("C", alloc_override=shift_from(BASE, "Inv", ["Liq", "Exist"], 10), label="inv_10"),
        "exist_5": run("C", alloc_override=reduce_exist_into_trav_inv(BASE, 5), label="exist_5"),
        "exist_10": run("C", alloc_override=reduce_exist_into_trav_inv(BASE, 10), label="exist_10"),
        "exist_15": run("C", alloc_override=reduce_exist_into_trav_inv(BASE, 15), label="exist_15"),
        "trf_follow": run("C", alloc_override=deepcopy(BASE), follow_trf=True, label="trf_follow"),
        "W": w,
    }

    rows_m = {k: metrics(v, r) for k, v in cases.items()}
    c_m = rows_m["C"]

    wb = Workbook()
    cover = wb.active
    cover.title = "00_Cover"
    cover["A1"] = "Phase 2.1 — use-of-funds incentives"
    cover["A1"].font = Font(bold=True, size=14)
    cover_lines = [
        "€700bn origination fixed. Same product as Central. Household cash-flow and outstanding unchanged.",
        "Government cannot assign household uses. These are tagged incentive shifts (Σ), not observed responses.",
        "Base mix: Diversifié-1. Renovation: 5/10/15 pp from Liq and Exist equally into Trav.",
        "Productive investment: 5/10 pp from Liq and Exist equally into Inv.",
        "Anti-asset-rotation: Exist −5/10/15 pp into Trav and Inv in the Diversifié-1 Trav:Inv ratio (28:5).",
        "Transfers followed: Trf euros are spent by the recipient on Diversifié-1 excluding nested transfers (Σ).",
        "Domestic transmission per €1 = additional French activity / €700bn. Not an established multiplier.",
        "Not wired: imports, labour/capacity vs prices, inflation, bank losses. No ranking. No EC-10. Not v0.1.",
    ]
    for i, line in enumerate(cover_lines, 3):
        cover.cell(i, 1, line)
    cover.column_dimensions["A"].width = 120

    alloc_rows = []
    labels = {
        "C": BASE,
        "reno_5": shift_from(BASE, "Trav", ["Liq", "Exist"], 5),
        "reno_10": shift_from(BASE, "Trav", ["Liq", "Exist"], 10),
        "reno_15": shift_from(BASE, "Trav", ["Liq", "Exist"], 15),
        "inv_5": shift_from(BASE, "Inv", ["Liq", "Exist"], 5),
        "inv_10": shift_from(BASE, "Inv", ["Liq", "Exist"], 10),
        "exist_5": reduce_exist_into_trav_inv(BASE, 5),
        "exist_10": reduce_exist_into_trav_inv(BASE, 10),
        "exist_15": reduce_exist_into_trav_inv(BASE, 15),
        "trf_follow": BASE,
        "W": ALLOC["Fuites-1"],
    }
    for name, alloc in labels.items():
        alloc_rows.append(
            [name]
            + [alloc[k] * 100 for k in USES]
            + [additional_spend(alloc) * 100]
            + ["yes" if name == "trf_follow" else "no"]
        )
    write_sheet(
        wb,
        "01_Allocations",
        ["case"] + USES + ["additional_spend_%", "follow_trf"],
        alloc_rows,
    )

    order = ["C", "reno_5", "reno_10", "reno_15", "inv_5", "inv_10", "exist_5", "exist_10", "exist_15", "trf_follow", "W"]
    write_sheet(
        wb,
        "02_Vs_R",
        [
            "case",
            "vol_bn",
            "rec_bn",
            "d_ratio_pp",
            "gross",
            "outstanding",
            "released",
            "vol_per_euro",
            "rec_per_euro",
        ],
        [
            [
                k,
                rows_m[k]["vol"],
                rows_m[k]["rec"],
                rows_m[k]["d_ratio_pp"],
                rows_m[k]["gross"],
                rows_m[k]["outstanding"],
                rows_m[k]["released"],
                rows_m[k]["vol_per_euro"],
                rows_m[k]["rec_per_euro"],
            ]
            for k in order
        ],
    )

    write_sheet(
        wb,
        "03_Delta_vs_C",
        ["case", "d_vol", "d_rec", "d_ratio_pp", "vol_per_euro", "rec_per_euro"],
        [
            [
                k,
                rows_m[k]["vol"] - c_m["vol"],
                rows_m[k]["rec"] - c_m["rec"],
                rows_m[k]["d_ratio_pp"] - c_m["d_ratio_pp"],
                rows_m[k]["vol_per_euro"],
                rows_m[k]["rec_per_euro"],
            ]
            for k in order
            if k != "C"
        ],
    )

    notes = wb.create_sheet("04_HowToRead")
    notes["A1"] = "How to read"
    notes["A1"].font = Font(bold=True, size=14)
    note_lines = [
        "€700bn is a constraint, not a target. Credit, cash-flow and outstanding do not move in this slice.",
        "C vs W is the envelope of the use-of-funds port, not a policy instrument.",
        "5/10/15 pp are Σ incentive intensities, not estimated elasticities of MaPrimeRénov’ or VAT.",
        "Anti-rotation landing in Trav/Inv is a policy-hope Σ. A conservative landing in Liq is not this run.",
        "Following transfers does not change the originator mix; it stops treating Trf as automatic leakage.",
        "vol_per_euro and rec_per_euro are ƒ of temporary EC-03/06 pass-throughs. Not established multipliers.",
        "Do not rank these interventions. Government assesses trade-offs. Other gates remain unwired.",
    ]
    for i, line in enumerate(note_lines, 3):
        notes.cell(i, 1, line)
    notes.column_dimensions["A"].width = 120

    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    print(f"Wrote {OUT}")
    print("--- vs R ---")
    for k in ["C", "W", "reno_10", "inv_10", "exist_10", "trf_follow"]:
        m = rows_m[k]
        print(
            f"{k}: vol={m['vol']:.1f} rec={m['rec']:.1f} d_ratio={m['d_ratio_pp']:.2f} "
            f"vol/€={m['vol_per_euro']:.3f} rec/€={m['rec_per_euro']:.3f} "
            f"gross={m['gross']:.0f} rel={m['released']:.1f} out={m['outstanding']:.0f}"
        )
    print("--- Δ vs C (hero intensities) ---")
    for k in ["reno_5", "reno_10", "reno_15", "inv_5", "inv_10", "exist_5", "exist_10", "exist_15", "trf_follow"]:
        m = rows_m[k]
        print(
            f"{k}: d_vol={m['vol']-c_m['vol']:.1f} d_rec={m['rec']-c_m['rec']:.1f} "
            f"d_ratio={m['d_ratio_pp']-c_m['d_ratio_pp']:.2f}"
        )


if __name__ == "__main__":
    main()
