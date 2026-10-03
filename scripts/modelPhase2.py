#!/usr/bin/env python3
"""Phase 2 tagged workbook. Regenerates Docs/20_Model_Phase2.xlsx.

Not a forecast. Not v0.1. Does not close any EC.
Tags: O observé · S sourcé · Σ scénario · ƒ calculé. C is the Central trajectory only.
"""

from __future__ import annotations

import math
from pathlib import Path

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
except ImportError as exc:  # pragma: no cover
    raise SystemExit("openpyxl is required: pip install openpyxl") from exc

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "Docs" / "20_Model_Phase2.xlsx"

YEARS = list(range(2025, 2041))
PACTE_YEARS = list(range(2027, 2041))

# EC-04 net-new IO path (Σ scenario column), Md€
NET_NEW = {
    2027: 25.0,
    2028: 45.0,
    2029: 65.0,
    2030: 80.0,
    2031: 90.0,
    2032: 90.0,
    2033: 80.0,
    2034: 70.0,
    2035: 55.0,
    2036: 40.0,
    2037: 30.0,
    2038: 20.0,
    2039: 10.0,
    2040: 0.0,
}

ALLOC = {
    "Diversifié-1": dict(Trav=0.28, Conso=0.12, Inv=0.05, Exist=0.15, Liq=0.15, Sub=0.15, Trf=0.10),
    "Fuites-1": dict(Trav=0.08, Conso=0.05, Inv=0.02, Exist=0.20, Liq=0.25, Sub=0.22, Trf=0.18),
    "Fuites-2": dict(Trav=0.05, Conso=0.04, Inv=0.01, Exist=0.30, Liq=0.22, Sub=0.20, Trf=0.18),
}

# Temporary Σ pass-throughs (Phase 2 visibility). Not EC-03 / EC-06 verdicts.
VOL = dict(Trav=0.40, goods=0.25, services=0.60, Inv=0.30)  # € volume FR / € spent
RECEIPT = dict(
    Trav=0.18,
    goods=0.17,
    services=0.25,
    Inv=0.15,
    Exist=0.06,  # DMTO-like on existing-asset euros
    Liq=0.0,
    Sub=0.0,
    Trf=0.0,
)

GDP_2025 = 2991.1  # O résumé / Insee
SPEND_RATIO_2025 = 0.572
DEFICIT_2025 = 152.5
DEBT_2025 = 3460.5
INTEREST_2025 = 64.7
U_2025 = 0.077  # O Insee BIT 2025 annual
DEFL = 0.025
G_REAL_R = 0.01
G_SPEND_EX = 0.02
REFIN_R = 0.0291
ROLL = 1.0 / (8.0 + 172.0 / 365.0)
HH_RATE_CW = 0.033  # Σ equal to amortising comparator (EC-05 WAL 3.30 %)
N_AMORT = 20
RECEIPT_ELASTICITY = 1.0  # Σ temporary vs nominal GDP
HH_GDI_SHARE = 0.70  # Σ share of GDP as disposable-income scale
DEFAULT_PER_PP = 0.001  # Σ 0.1 % of outstanding per pp unemployment gap


def annuity_payment(principal: float, rate: float, n: int) -> float:
    if principal <= 0:
        return 0.0
    if rate <= 0:
        return principal / n
    return principal * rate / (1.0 - (1.0 + rate) ** (-n))


def cash_released_on_vintage(principal: float, rate: float) -> float:
    return annuity_payment(principal, rate, N_AMORT) - rate * principal


def hp_index(traj: str, year: int) -> float:
    """2027 = 100. R follows deflator. Adverse housing: 70 by 2030, then held at 70."""
    if year < 2027:
        return 100.0 / ((1.0 + DEFL) ** (2027 - year))
    housing_shock = traj in {
        "A0",
        "A",
        "A_no_recession",
        "A_no_rates",
        "A_Fuites-1",
        "A_no_stop",
        "A_stop_now",
    }
    if not housing_shock:
        return 100.0 * ((1.0 + DEFL) ** (year - 2027))
    if year <= 2027:
        return 100.0
    if year == 2028:
        return 90.0
    if year == 2029:
        return 80.0
    return 70.0


def real_growth(traj: str, year: int) -> float:
    if year <= 2025:
        return 0.0
    recession = traj in {
        "A0",
        "A",
        "A_no_housing",
        "A_no_rates",
        "A_Fuites-1",
        "A_no_stop",
        "A_stop_now",
    }
    if recession and year in (2028, 2029, 2030):
        return G_REAL_R - 0.015
    return G_REAL_R


def unemployment(traj: str, year: int) -> float:
    shock = traj in {
        "A0",
        "A",
        "A_no_housing",
        "A_no_rates",
        "A_Fuites-1",
        "A_no_stop",
        "A_stop_now",
    }
    if shock and year in (2028, 2029, 2030):
        return U_2025 + 0.02
    return U_2025


def refin_rate(traj: str, year: int) -> float:
    high = traj in {
        "A0",
        "A",
        "A_no_recession",
        "A_no_housing",
        "A_Fuites-1",
        "A_no_stop",
        "A_stop_now",
    }
    if high and year >= 2028:
        return REFIN_R + 0.02
    return REFIN_R


def hh_new_rate(traj: str, year: int) -> float:
    high = traj in {
        "A",
        "A_no_recession",
        "A_no_housing",
        "A_Fuites-1",
        "A_no_stop",
        "A_stop_now",
    }
    if high and year >= 2028:
        return HH_RATE_CW + 0.02
    return HH_RATE_CW


def allocation_name(traj: str) -> str | None:
    if traj in {"R", "A0"}:
        return None
    if traj == "C":
        return "Diversifié-1"
    if traj == "W":
        return "Fuites-1"
    if traj == "A_Fuites-1":
        return "Fuites-1"
    return "Fuites-2"


def gross_for_year(traj: str, year: int) -> float:
    if traj in {"R", "A0"} or year not in NET_NEW:
        return 0.0
    scheduled = NET_NEW[year]
    if traj == "A_stop_now" and year >= 2028:
        return 0.0
    if traj in {"A", "A_no_recession", "A_no_housing", "A_no_rates", "A_Fuites-1"} and year >= 2029:
        return 0.0
    return scheduled


def mix_without_trf(alloc: dict) -> dict:
    rest = {k: v for k, v in alloc.items() if k != "Trf"}
    total = sum(rest.values())
    if total <= 0:
        return dict(alloc)
    return {k: v / total for k, v in rest.items()} | {"Trf": 0.0}


def volume_and_receipts(
    gross: float,
    alloc: dict | None,
    hp: float | None,
    follow_trf: bool = False,
) -> tuple[float, float]:
    if not alloc or gross <= 0:
        return 0.0, 0.0
    vol = gross * (
        alloc["Trav"] * VOL["Trav"]
        + alloc["Conso"] * 0.5 * VOL["goods"]
        + alloc["Conso"] * 0.5 * VOL["services"]
        + alloc["Inv"] * VOL["Inv"]
    )
    hp_scale = (hp / 100.0) if hp else 1.0
    rec = gross * (
        alloc["Trav"] * RECEIPT["Trav"]
        + alloc["Conso"] * 0.5 * RECEIPT["goods"]
        + alloc["Conso"] * 0.5 * RECEIPT["services"]
        + alloc["Inv"] * RECEIPT["Inv"]
        + alloc["Exist"] * RECEIPT["Exist"] * hp_scale
        + alloc["Liq"] * RECEIPT["Liq"]
        + alloc["Sub"] * RECEIPT["Sub"]
        + alloc["Trf"] * RECEIPT["Trf"]
    )
    if follow_trf and alloc.get("Trf", 0) > 0:
        nested_vol, nested_rec = volume_and_receipts(
            gross * alloc["Trf"],
            mix_without_trf(alloc),
            hp,
            follow_trf=False,
        )
        vol += nested_vol
        rec += nested_rec
    return vol, rec


def run(
    traj: str,
    *,
    alloc_override: dict | None = None,
    follow_trf: bool = False,
    label: str | None = None,
) -> dict:
    spend_2025 = SPEND_RATIO_2025 * GDP_2025
    receipts_2025 = spend_2025 - DEFICIT_2025
    spend_ex = spend_2025 - INTEREST_2025
    eff_rate = INTEREST_2025 / DEBT_2025

    price = 1.0
    real = GDP_2025
    nom = GDP_2025
    debt = DEBT_2025
    rec_base = receipts_2025
    vintages: list[tuple[int, float, float]] = []  # year, principal still out, rate

    rows = []
    for year in YEARS:
        g = real_growth(traj, year) if year > 2025 else 0.0
        if year > 2025:
            price *= 1.0 + DEFL
            real *= 1.0 + g
            spend_ex *= 1.0 + G_SPEND_EX
            g_nom_base = (1.0 + g) * (1.0 + DEFL) - 1.0
            rec_base *= 1.0 + g_nom_base * RECEIPT_ELASTICITY

        gross = gross_for_year(traj, year)
        repay = 0.0
        alloc = (
            alloc_override
            if alloc_override is not None
            else ALLOC.get(allocation_name(traj) or "", None)
        )
        hp = hp_index(traj, year) if year >= 2027 else 100.0
        vol, pact_rec = volume_and_receipts(
            gross, alloc, hp if year >= 2027 else None, follow_trf=follow_trf
        )
        if year > 2025:
            real += vol / price
            nom = real * price
        else:
            nom = GDP_2025

        if gross > 0:
            vintages.append((year, gross, hh_new_rate(traj, year)))
        outstanding = sum(p for _, p, _ in vintages)
        hh_interest = sum(p * r for _, p, r in vintages)
        released = sum(cash_released_on_vintage(p, r) for _, p, r in vintages)
        u = unemployment(traj, year)
        defaults = outstanding * DEFAULT_PER_PP * max(0.0, (u - U_2025) * 100.0)

        if year > 2025:
            ref = refin_rate(traj, year)
            eff_rate = (1.0 - ROLL) * eff_rate + ROLL * ref
            interest = eff_rate * debt
            receipts = rec_base + pact_rec
            spend_tot = spend_ex + interest
            deficit = spend_tot - receipts
            primary = (receipts - spend_ex)
            debt = debt + deficit
        else:
            interest = INTEREST_2025
            receipts = receipts_2025
            spend_tot = spend_2025
            deficit = DEFICIT_2025
            primary = receipts_2025 - spend_ex
            debt = DEBT_2025

        spread = ((nom / (rows[-1]["nom"] if rows else nom)) - 1.0) - G_SPEND_EX if rows else None
        # spend growth is on ex-interest; Spread in V2 is vs total public spending nominal.
        # Use total spending growth vs nominal GDP growth when prior year exists.
        if rows:
            g_nom = nom / rows[-1]["nom"] - 1.0
            g_sp = spend_tot / rows[-1]["spend_tot"] - 1.0
            spread = g_nom - g_sp
        else:
            spread = None

        rows.append(
            dict(
                year=year,
                g_real=g if year > 2025 else None,
                u=u,
                hp=hp if year >= 2027 else None,
                refin=refin_rate(traj, year) if year > 2025 else None,
                gross=gross,
                repay=repay,
                outstanding=outstanding,
                cumul=sum(NET_NEW[y] for y in PACTE_YEARS if y <= year)
                if traj not in {"R"}
                else 0.0,
                vol=vol,
                pact_rec=pact_rec,
                real=real,
                nom=nom,
                receipts=receipts,
                spend_ex=spend_ex,
                spend_tot=spend_tot,
                interest=interest,
                primary=primary,
                deficit=deficit,
                debt=debt,
                ratio=debt / nom,
                spread=spread,
                hh_interest=hh_interest,
                released=released,
                defaults=defaults,
                gdi=HH_GDI_SHARE * nom,
            )
        )
        # fix cumulative origination to actual gross sum not scheduled path
        rows[-1]["cumul"] = (rows[-2]["cumul"] if len(rows) > 1 else 0.0) + gross

    return {
        "traj": label or traj,
        "rows": rows,
        "alloc": alloc_override if alloc_override is not None else allocation_name(traj),
        "follow_trf": follow_trf,
    }


TRAJECTORIES = ["R", "C", "W", "A0", "A"]
ABLATIONS = ["A_no_recession", "A_no_housing", "A_no_rates", "A_Fuites-1", "A_no_stop", "A_stop_now"]


def row2040(run_out: dict) -> dict:
    return next(r for r in run_out["rows"] if r["year"] == 2040)


def sum_field(run_out: dict, field: str, from_year: int = 2027) -> float:
    return sum(r[field] for r in run_out["rows"] if r["year"] >= from_year)


def dominance_block(runs: dict) -> list[dict]:
    r = runs["R"]
    r40 = row2040(r)
    out = []
    for name in ["C", "W", "A0", "A"]:
        x = runs[name]
        x40 = row2040(x)
        out.append(
            dict(
                name=name,
                vol_cum=sum_field(x, "vol") - sum_field(r, "vol"),
                rec_cum=sum_field(x, "pact_rec") - sum_field(r, "pact_rec"),
                d_primary_2040=x40["primary"] - r40["primary"],
                d_interest_2040=x40["interest"] - r40["interest"],
                d_deficit_2040=x40["deficit"] - r40["deficit"],
                d_debt_2040=x40["debt"] - r40["debt"],
                d_ratio_pp=100 * (x40["ratio"] - r40["ratio"]),
                outstanding_2040=x40["outstanding"],
                released_cum=sum_field(x, "released"),
                hh_int_2040=x40["hh_interest"],
                service_income_2040=x40["hh_interest"] / x40["gdi"] if x40["gdi"] else None,
                defaults_cum=sum_field(x, "defaults"),
                gross_cum=x40["cumul"],
            )
        )
    return out


def ablation_block(runs: dict) -> list[dict]:
    a = runs["A"]
    a40 = row2040(a)
    out = []
    for name in ABLATIONS:
        x = runs[name]
        x40 = row2040(x)
        out.append(
            dict(
                name=name,
                d_vol=sum_field(x, "vol") - sum_field(a, "vol"),
                d_rec=sum_field(x, "pact_rec") - sum_field(a, "pact_rec"),
                d_ratio_pp=100 * (x40["ratio"] - a40["ratio"]),
                d_interest=x40["interest"] - a40["interest"],
                d_out=x40["outstanding"] - a40["outstanding"],
                d_released=sum_field(x, "released") - sum_field(a, "released"),
                d_hh_int=x40["hh_interest"] - a40["hh_interest"],
            )
        )
    return out


FILL_O = PatternFill("solid", fgColor="D9EAD3")
FILL_S = PatternFill("solid", fgColor="FFF2CC")
FILL_SIG = PatternFill("solid", fgColor="FCE5CD")
FILL_F = PatternFill("solid", fgColor="D0E2F3")
HEAD = Font(bold=True)


def autosize(ws) -> None:
    for col in ws.columns:
        letter = get_column_letter(col[0].column)
        width = 12
        for cell in col:
            if cell.value is None:
                continue
            width = min(48, max(width, len(str(cell.value)) + 2))
        ws.column_dimensions[letter].width = width


def write_sheet(wb: Workbook, title: str, headers: list[str], rows: list[list], fills: dict | None = None):
    ws = wb.create_sheet(title)
    for j, h in enumerate(headers, 1):
        c = ws.cell(1, j, h)
        c.font = HEAD
    for i, row in enumerate(rows, 2):
        for j, val in enumerate(row, 1):
            cell = ws.cell(i, j, val)
            if fills and headers[j - 1] in fills:
                cell.fill = fills[headers[j - 1]]
            if isinstance(val, float):
                cell.number_format = "0.000"
                cell.alignment = Alignment(horizontal="right")
    autosize(ws)
    return ws


def main() -> None:
    runs = {name: run(name) for name in TRAJECTORIES + ABLATIONS}
    dom = dominance_block(runs)
    abl = ablation_block(runs)

    wb = Workbook()
    cover = wb.active
    cover.title = "00_Readme"
    cover_lines = [
        "France 2040 — Phase 2 tagged model",
        "The model does not resolve uncertainty. It makes uncertainty compete.",
        "Not a forecast. Not v0.1. Does not close any EC. Not a public snapshot (Phase 3).",
        "Tags: O observé · S sourcé · Σ scénario · ƒ calculé. C = Central trajectory only.",
        "Rebuild: python3 scripts/modelPhase2.py",
        "",
        "R = France without Pacte (tagged counterfactual).",
        "C = frozen net IO path + Diversifié-1.",
        "W = same path + Fuites-1.",
        "A0 = V2 adverse macro, no Pacte (adverse France without the mechanism).",
        "A = A0 + Fuites-2 + first light 2028 + stop from 2029.",
        "Housing A: index 2027=100, 70 in 2030, held at 70 (nominal −30 %, not 30 % below R).",
        "Bullets: 2027 vintage matures 2047; each vintage +20 years. None in 2027–2040.",
        "Stop clock (Σ, this run): first light = 2028 recession year; delay 4 quarters → stop 2029+.",
        "EC-03 / EC-06 coefficients are temporary Σ pass-throughs on sheet 03_PassThroughs.",
        "Ablations remove one A component at a time (sheet 09_Ablations).",
    ]
    for i, line in enumerate(cover_lines, 1):
        cover.cell(i, 1, line)
        if i == 2:
            cover.cell(i, 1).font = Font(italic=True)
        if i == 1:
            cover.cell(i, 1).font = Font(bold=True, size=14)
    cover.column_dimensions["A"].width = 110

    write_sheet(
        wb,
        "01_Assumptions",
        ["cell", "tag", "value", "justification"],
        [
            ["GDP 2025, Md€", "O", GDP_2025, "Résumé / Insee; start not 2040"],
            ["Public spending 2025 / GDP", "O", SPEND_RATIO_2025, "Insee IR 78: 57.2 %"],
            ["Deficit 2025 Md€", "O", DEFICIT_2025, "Insee Première 2106"],
            ["Debt 2025 Md€", "O", DEBT_2025, "Insee Première 2106 Maastricht"],
            ["Interest 2025 Md€", "O", INTEREST_2025, "Insee 64.7"],
            ["Unemployment BIT 2025", "O", U_2025, "Insee enquête Emploi 2025 annual 7.7 %"],
            ["Real growth R", "Σ", G_REAL_R, "Résumé working saisie; not Insee potential"],
            ["GDP deflator", "Σ", DEFL, "Résumé working saisie"],
            ["Spend ex-interest growth", "Σ", G_SPEND_EX, "Résumé 2.0 %; interest is separate ƒ"],
            ["Refinancing rate R", "Σ", REFIN_R, "Hold AFT 2024 MLT 2.91 % (O level, Σ path)"],
            ["Roll fraction of public debt", "S", ROLL, "1 / (8 years + 172 days), AFT 2024"],
            ["HH IO rate C/W new origination", "Σ", HH_RATE_CW, "Set equal to amortising comparator so cash released is amortisation avoided, not a rate gift"],
            ["Comparator amortising rate / maturity", "Σ", "3.30 % / 20y", "EC-05 WAL arithmetic"],
            ["Prepayment 2027–2040", "Σ", 0, "Do not import SFH CPR onto the Pacte"],
            ["Receipt elasticity vs nominal GDP", "Σ", RECEIPT_ELASTICITY, "Temporary Phase 2; not 43.6 % marginal"],
            ["A recession", "Σ", "R−1.5pp in 2028–2030", "V2 recession, round gap"],
            ["A unemployment", "Σ", "R+2pp in 2028–2030", "EC-09 strand"],
            ["A house-price index", "Σ", "100 in 2027 → 70 in 2030, then 70", "Nominal −30 % from start; not 30 % below R"],
            ["A sovereign / HH new rate", "Σ", "+200 bp from 2028", "V2 high rates; existing vintages stay fixed"],
            ["A first light", "Σ", 2028, "This-run clock = first recession year; not a policy threshold"],
            ["A stop", "Σ", "from 2029 inclusive", "4 quarters after 2028"],
            ["Stock-flow on debt", "Σ", "debt(t)=debt(t-1)+deficit(t)", "Ignores 2025 cash adjustment; EC-07 identity simplified"],
        ],
    )

    alloc_rows = []
    for name, vec in ALLOC.items():
        s = 0.0
        for k, v in vec.items():
            alloc_rows.append([name, k, "Σ", v])
            s += v
        alloc_rows.append([name, "additional_spend ƒ", "ƒ", vec["Trav"] + vec["Conso"] + vec["Inv"]])
        alloc_rows.append([name, "sum", "ƒ", s])
    write_sheet(wb, "02_Allocations", ["vector", "use", "tag", "share"], alloc_rows)

    write_sheet(
        wb,
        "03_PassThroughs",
        ["port", "tag", "value", "justification"],
        [
            ["VOL Trav", "Σ", VOL["Trav"], "Temporary: high FR content but rigid supply → half-ish volume not 96 %"],
            ["VOL Conso goods", "Σ", VOL["goods"], "Temporary: below 38 % average content, visibility for imports"],
            ["VOL Conso services", "Σ", VOL["services"], "Temporary: local labour; not 80 % as a Pacte param"],
            ["VOL Inv", "Σ", VOL["Inv"], "Temporary: equipment often imported"],
            ["RECEIPT Trav", "Σ", RECEIPT["Trav"], "Temporary: reduced VAT + some labour if volume; not 43.6 %"],
            ["RECEIPT goods", "Σ", RECEIPT["goods"], "Temporary: ~VAT 20 % on TTC"],
            ["RECEIPT services", "Σ", RECEIPT["services"], "Temporary: VAT + labour"],
            ["RECEIPT Inv", "Σ", RECEIPT["Inv"], "Temporary"],
            ["RECEIPT Exist", "Σ", RECEIPT["Exist"], "Temporary DMTO-like on existing-asset euros; Fuites-2 wedge lives here"],
            ["Conso split goods/services", "Σ", 0.5, "No observed Pacte basket"],
        ],
    )

    cred_headers = ["traj", "year", "tag", "gross", "repay", "cumul_orig", "outstanding", "net_new_if_repay0"]
    cred_rows = []
    for name in TRAJECTORIES:
        for r in runs[name]["rows"]:
            if r["year"] < 2027:
                continue
            cred_rows.append(
                [name, r["year"], "ƒ", r["gross"], r["repay"], r["cumul"], r["outstanding"], NET_NEW.get(r["year"], 0.0) if name != "R" else 0.0]
            )
    write_sheet(wb, "04_Credit", cred_headers, cred_rows)

    hh_rows = []
    for name in TRAJECTORIES:
        for r in runs[name]["rows"]:
            if r["year"] < 2027:
                continue
            hh_rows.append(
                [
                    name,
                    r["year"],
                    r["outstanding"],
                    r["released"],
                    r["hh_interest"],
                    r["gdi"],
                    (r["hh_interest"] / r["gdi"]) if r["gdi"] else None,
                    r["defaults"],
                    r["hp"],
                    r["u"],
                ]
            )
    write_sheet(
        wb,
        "05_Household",
        [
            "traj",
            "year",
            "outstanding_IO",
            "cash_released_amort_avoided",
            "IO_interest",
            "GDI_scale",
            "interest/GDI",
            "defaults_Σ",
            "house_price_index",
            "unemployment",
        ],
        hh_rows,
    )

    mac_rows = []
    for name in TRAJECTORIES:
        for r in runs[name]["rows"]:
            mac_rows.append(
                [
                    name,
                    r["year"],
                    r["g_real"],
                    r["real"],
                    r["nom"],
                    r["vol"],
                    r["pact_rec"],
                    r["receipts"],
                    r["spend_ex"],
                    r["interest"],
                    r["spend_tot"],
                    r["primary"],
                    r["deficit"],
                    r["debt"],
                    r["ratio"],
                    r["spread"],
                    r["refin"],
                ]
            )
    write_sheet(
        wb,
        "06_MacroFinance",
        [
            "traj",
            "year",
            "g_real",
            "real_GDP",
            "nom_GDP",
            "FR_volume_impulse",
            "pacte_receipts",
            "receipts",
            "spend_ex_interest",
            "public_interest",
            "spend_tot",
            "primary",
            "deficit",
            "debt",
            "debt/GDP",
            "Spread",
            "refin_rate",
        ],
        mac_rows,
    )

    write_sheet(
        wb,
        "08_Dominance",
        [
            "vs R",
            "FR volume 2027–40",
            "Pacte receipts 2027–40",
            "Δ primary 2040",
            "Δ public interest 2040",
            "Δ deficit 2040",
            "Δ debt 2040",
            "Δ ratio pp 2040",
            "IO outstanding 2040",
            "cash released 2027–40",
            "HH IO interest 2040",
            "HH interest/GDI 2040",
            "defaults 2027–40",
            "gross orig 2027–40",
        ],
        [
            [
                d["name"],
                d["vol_cum"],
                d["rec_cum"],
                d["d_primary_2040"],
                d["d_interest_2040"],
                d["d_deficit_2040"],
                d["d_debt_2040"],
                d["d_ratio_pp"],
                d["outstanding_2040"],
                d["released_cum"],
                d["hh_int_2040"],
                d["service_income_2040"],
                d["defaults_cum"],
                d["gross_cum"],
            ]
            for d in dom
        ],
    )

    write_sheet(
        wb,
        "09_Ablations",
        [
            "vs A",
            "meaning",
            "Δ FR volume",
            "Δ Pacte receipts",
            "Δ ratio pp 2040",
            "Δ public interest 2040",
            "Δ outstanding 2040",
            "Δ cash released",
            "Δ HH interest 2040",
        ],
        [
            [
                a["name"],
                {
                    "A_no_recession": "keep Fuites-2, housing, rates, stop; drop growth/unemployment gap",
                    "A_no_housing": "index follows R deflator (no 70-by-2030)",
                    "A_no_rates": "no +200 bp",
                    "A_Fuites-1": "swap allocation to Fuites-1 (drop Exist wedge)",
                    "A_no_stop": "full EC-04 path (delay never binds)",
                    "A_stop_now": "stop from 2028 (delay 0)",
                }[a["name"]],
                a["d_vol"],
                a["d_rec"],
                a["d_ratio_pp"],
                a["d_interest"],
                a["d_out"],
                a["d_released"],
                a["d_hh_int"],
            ]
            for a in abl
        ],
    )

    # Read-me numbers for the note
    notes = wb.create_sheet("10_HowToRead")
    notes["A1"] = "How to read dominance"
    notes["A1"].font = Font(bold=True, size=14)
    lines = [
        "C − R is not a proof of the thesis. It is Diversifié-1 plus temporary pass-throughs competing against R’s tagged saisies.",
        "W − R shows how much of C’s impulse dies when the allocation is Fuites-1 (same credit path).",
        "A0 is the adverse world without the Pacte. Compare A to A0, not only A to R.",
        "C − R is the expected central mechanism, not a proof of the thesis.",
        "If A_no_stop moves outstanding and household service far more than A_Fuites-1, the stop convention dominates the Exist wedge.",
        "If A_Fuites-1 barely moves volume/receipts vs A, the Exist 30 % wedge is small in this pass-through — deepen EC-06 only if receipts/DMTO move.",
        "If C − R on the ratio is mostly R’s 2.5 % deflator + constant-spend 2 % (look at R alone), the Pacte overlay is not doing the work.",
        "700 cumulative = 700 outstanding in C/W only because Σ prepayment = 0. That is ƒ, not an identity of the Pacte.",
        "First bullets: 2047 for 2027 vintage, then vintage+20. None in this window.",
    ]
    for i, line in enumerate(lines, 3):
        notes.cell(i, 1, line)
    notes.column_dimensions["A"].width = 120

    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)

    # stdout summary for the chat / Docs/20 markdown writer
    print(f"Wrote {OUT}")
    print("--- Dominance vs R (2040 / cumul) ---")
    for d in dom:
        print(
            f"{d['name']}: vol={d['vol_cum']:.1f} rec={d['rec_cum']:.1f} "
            f"d_ratio_pp={d['d_ratio_pp']:.2f} out={d['outstanding_2040']:.1f} "
            f"released={d['released_cum']:.1f} hh_int={d['hh_int_2040']:.1f} gross={d['gross_cum']:.1f}"
        )
    print("--- A vs A0 (Pacte inside the adverse world) ---")
    a0 = runs["A0"]
    a = runs["A"]
    a040 = row2040(a0)
    a40 = row2040(a)
    r40 = row2040(runs["R"])
    print(
        f"A0 vs R ratio_pp={100 * (a040['ratio'] - r40['ratio']):.2f} "
        f"A vs A0 ratio_pp={100 * (a40['ratio'] - a040['ratio']):.2f} "
        f"A vs R ratio_pp={100 * (a40['ratio'] - r40['ratio']):.2f}"
    )
    print(
        f"A vs A0: d_vol={sum_field(a, 'vol') - sum_field(a0, 'vol'):.1f} "
        f"d_rec={sum_field(a, 'pact_rec') - sum_field(a0, 'pact_rec'):.1f} "
        f"d_interest={a40['interest'] - a040['interest']:.1f} "
        f"out={a40['outstanding']:.1f}"
    )
    print("--- Ablations vs A ---")
    for a in abl:
        print(
            f"{a['name']}: d_vol={a['d_vol']:.2f} d_rec={a['d_rec']:.2f} "
            f"d_ratio_pp={a['d_ratio_pp']:.2f} d_out={a['d_out']:.1f} d_hh={a['d_hh_int']:.2f}"
        )


if __name__ == "__main__":
    main()
